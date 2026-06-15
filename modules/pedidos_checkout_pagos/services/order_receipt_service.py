from __future__ import annotations

from dataclasses import dataclass
from io import BytesIO
from uuid import UUID
import html

from django.db import transaction
from django.utils import timezone

from modules.gestion_cocinero.models import ChefProfile
from modules.gestion_usuarios_acceso_suscripcion.models import UserProfile
from modules.pedidos_checkout_pagos.models import Order, OrderPayment, OrderReceipt, OrderTimelineEvent


@dataclass
class ReceiptDownloadPayload:
    content: bytes
    content_type: str
    filename: str


class OrderReceiptServiceError(ValueError):
    def __init__(self, message: str, code: str, details: dict | None = None):
        super().__init__(message)
        self.message = message
        self.code = code
        self.details = details or {}


class OrderReceiptService:
    FORMAT_CONTENT_TYPES = {
        "pdf": "application/pdf",
        "html": "text/html; charset=utf-8",
        "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    }

    @transaction.atomic
    def ensure_receipt(
        self,
        order: Order,
        payment: OrderPayment | None,
        *,
        actor_role: str = "SISTEMA",
        actor_id: str = "",
        metadata: dict | None = None,
    ) -> OrderReceipt | None:
        if not payment or payment.status != OrderPayment.Status.CONFIRMED:
            return None

        defaults = self._build_snapshot(order, payment, metadata)
        receipt, created = OrderReceipt.objects.get_or_create(
            order=order,
            payment=payment,
            defaults={
                "receipt_number": self._receipt_number(order, payment),
                **defaults,
            },
        )
        dirty_fields = []
        for field, value in defaults.items():
            if getattr(receipt, field) != value:
                setattr(receipt, field, value)
                dirty_fields.append(field)
        if dirty_fields:
            dirty_fields.extend(["updated_at"])
            receipt.save(update_fields=dirty_fields)
        if created:
            OrderTimelineEvent.objects.create(
                order=order,
                event_code="PAYMENT_RECEIPT_GENERATED",
                event_label="Comprobante de pago generado",
                actor_role=actor_role,
                actor_id=str(actor_id),
                metadata={
                    "receipt_id": receipt.id,
                    "receipt_number": receipt.receipt_number,
                    "payment_id": payment.id,
                    "payment_method": payment.method,
                },
            )
        return receipt

    def list_client_receipts(self, user_id: str, order_id: str):
        client = self._require_profile(user_id, UserProfile.ROLE_CLIENT, "client_not_found", "Perfil de cliente no encontrado.")
        order = self._get_order_for_client(client, order_id)
        return {"items": self._serialize_receipts(order)}

    def list_chef_receipts(self, user_id: str, order_id: str):
        chef = self._require_profile(user_id, UserProfile.ROLE_CHEF, "chef_not_found", "Perfil de cocinero no encontrado.")
        order = self._get_order_for_chef(chef, order_id)
        return {"items": self._serialize_receipts(order)}

    @transaction.atomic
    def download_client_receipt(self, user_id: str, order_id: str, receipt_id: str, format_value: str):
        client = self._require_profile(user_id, UserProfile.ROLE_CLIENT, "client_not_found", "Perfil de cliente no encontrado.")
        order = self._get_order_for_client(client, order_id)
        return self._download(order, receipt_id, format_value, actor_role="CLIENTE", actor_id=str(client.supabase_user_id))

    @transaction.atomic
    def download_chef_receipt(self, user_id: str, order_id: str, receipt_id: str, format_value: str):
        chef = self._require_profile(user_id, UserProfile.ROLE_CHEF, "chef_not_found", "Perfil de cocinero no encontrado.")
        order = self._get_order_for_chef(chef, order_id)
        return self._download(order, receipt_id, format_value, actor_role="COCINERO", actor_id=str(chef.supabase_user_id))

    def serialize_receipts_for_order(self, order: Order):
        self._ensure_receipt_for_confirmed_payment(order)
        return self._serialize_receipts(order)

    def _download(self, order: Order, receipt_id: str, format_value: str, *, actor_role: str, actor_id: str):
        normalized_format = str(format_value or "pdf").lower()
        if normalized_format not in self.FORMAT_CONTENT_TYPES:
            raise OrderReceiptServiceError("Formato de comprobante no soportado.", "receipt_format_invalid")
        receipt = order.receipts.filter(id=str(receipt_id)).select_related("payment").first()
        if not receipt:
            raise OrderReceiptServiceError("Comprobante no encontrado para el pedido.", "receipt_not_found")
        payload = self._build_download_payload(order, receipt, normalized_format)
        OrderTimelineEvent.objects.create(
            order=order,
            event_code="PAYMENT_RECEIPT_DOWNLOADED",
            event_label="Comprobante descargado",
            actor_role=actor_role,
            actor_id=str(actor_id),
            metadata={
                "receipt_id": receipt.id,
                "receipt_number": receipt.receipt_number,
                "format": normalized_format,
            },
        )
        return payload

    def _build_download_payload(self, order: Order, receipt: OrderReceipt, format_value: str):
        if format_value == "html":
            content = self._render_html(order, receipt).encode("utf-8")
        elif format_value == "docx":
            content = self._render_docx(order, receipt)
        else:
            content = self._render_pdf(order, receipt)
        return ReceiptDownloadPayload(
            content=content,
            content_type=self.FORMAT_CONTENT_TYPES[format_value],
            filename=f"{receipt.receipt_number}.{format_value}",
        )

    def _serialize_receipts(self, order: Order):
        self._ensure_receipt_for_confirmed_payment(order)
        receipts = order.receipts.all().order_by("-issued_at")
        return [
            {
                "id": receipt.id,
                "receipt_number": receipt.receipt_number,
                "payment_method": receipt.payment_method,
                "payment_status": receipt.payment_status,
                "order_status": receipt.order_status,
                "currency": receipt.currency,
                "total": float(receipt.total),
                "external_reference": receipt.external_reference,
                "issued_at": receipt.issued_at.isoformat(),
                "confirmed_at": receipt.confirmed_at.isoformat() if receipt.confirmed_at else None,
                "available_formats": ["pdf", "html", "docx"],
            }
            for receipt in receipts
        ]

    def _ensure_receipt_for_confirmed_payment(self, order: Order):
        payment = order.payments.order_by("-created_at").first()
        if not payment or payment.status != OrderPayment.Status.CONFIRMED:
            return None
        if order.receipts.filter(payment=payment).exists():
            return None
        return self.ensure_receipt(order, payment)

    def _build_snapshot(self, order: Order, payment: OrderPayment, metadata: dict | None):
        return {
            "payment_method": payment.method,
            "payment_status": payment.status,
            "order_status": order.status,
            "currency": order.currency,
            "subtotal": order.subtotal,
            "delivery_fee": order.delivery_fee,
            "service_fee": order.service_fee,
            "discount_total": order.discount_total,
            "total": order.total,
            "external_reference": payment.external_reference,
            "confirmed_at": payment.confirmed_at,
            "metadata": {
                **dict(payment.provider_payload or {}),
                **dict(metadata or {}),
            },
        }

    def _receipt_number(self, order: Order, payment: OrderPayment):
        timestamp = timezone.now().strftime("%Y%m%d%H%M%S")
        return f"HC-{timestamp}-{str(order.id)[:8].upper()}-{str(payment.id)[:6].upper()}"

    def _render_html(self, order: Order, receipt: OrderReceipt):
        context = self._context(order, receipt)
        rows = "".join(
            f"""
            <tr>
              <td>{html.escape(item['dish_name'])}</td>
              <td>{item['quantity']}</td>
              <td>{item['unit_price']}</td>
              <td>{item['subtotal']}</td>
            </tr>
            """
            for item in context["items"]
        )
        return f"""<!doctype html>
<html lang="es">
  <head>
    <meta charset="utf-8" />
    <title>Comprobante {html.escape(context['receipt_number'])}</title>
    <style>
      body {{ font-family: Arial, sans-serif; color: #0f172a; margin: 32px; }}
      .shell {{ max-width: 860px; margin: 0 auto; border: 1px solid #cbd5e1; border-radius: 20px; padding: 28px; }}
      .head {{ display: flex; justify-content: space-between; gap: 24px; margin-bottom: 24px; }}
      .brand {{ font-size: 28px; font-weight: 700; }}
      .muted {{ color: #475569; }}
      .grid {{ display: grid; grid-template-columns: repeat(2, minmax(0, 1fr)); gap: 12px; margin-bottom: 24px; }}
      .card {{ border: 1px solid #e2e8f0; border-radius: 14px; padding: 14px; }}
      table {{ width: 100%; border-collapse: collapse; margin-top: 16px; }}
      th, td {{ border-bottom: 1px solid #e2e8f0; padding: 10px 8px; text-align: left; }}
      .totals {{ margin-top: 18px; width: 100%; max-width: 320px; margin-left: auto; }}
      .totals td {{ border: none; padding: 6px 0; }}
      .footer {{ margin-top: 28px; font-size: 13px; color: #475569; }}
    </style>
  </head>
  <body>
    <div class="shell">
      <div class="head">
        <div>
          <div class="brand">HomeChef</div>
          <div class="muted">Comprobante formal de pago</div>
        </div>
        <div>
          <div><strong>Nro:</strong> {html.escape(context['receipt_number'])}</div>
          <div><strong>Emitido:</strong> {html.escape(context['issued_at'])}</div>
          <div><strong>Pedido:</strong> {html.escape(context['order_id'])}</div>
        </div>
      </div>
      <div class="grid">
        <div class="card">
          <strong>Cliente</strong>
          <div>{html.escape(context['client_name'])}</div>
        </div>
        <div class="card">
          <strong>Cocinero</strong>
          <div>{html.escape(context['chef_name'])}</div>
        </div>
        <div class="card">
          <strong>Modalidad</strong>
          <div>{html.escape(context['fulfillment_type'])}</div>
        </div>
        <div class="card">
          <strong>Metodo de pago</strong>
          <div>{html.escape(context['payment_method'])}</div>
        </div>
        <div class="card">
          <strong>Estado del pago</strong>
          <div>{html.escape(context['payment_status'])}</div>
        </div>
        <div class="card">
          <strong>Referencia externa</strong>
          <div>{html.escape(context['external_reference'])}</div>
        </div>
      </div>
      <table>
        <thead>
          <tr>
            <th>Plato</th>
            <th>Cantidad</th>
            <th>Unitario</th>
            <th>Subtotal</th>
          </tr>
        </thead>
        <tbody>{rows}</tbody>
      </table>
      <table class="totals">
        <tr><td><strong>Subtotal</strong></td><td>{context['subtotal']}</td></tr>
        <tr><td><strong>Delivery</strong></td><td>{context['delivery_fee']}</td></tr>
        <tr><td><strong>Servicio</strong></td><td>{context['service_fee']}</td></tr>
        <tr><td><strong>Descuento</strong></td><td>{context['discount_total']}</td></tr>
        <tr><td><strong>Total</strong></td><td><strong>{context['total']}</strong></td></tr>
      </table>
      <div class="footer">
        Estado final del pedido al emitir: {html.escape(context['order_status'])}. Direccion: {html.escape(context['address'])}
      </div>
    </div>
  </body>
</html>"""

    def _render_docx(self, order: Order, receipt: OrderReceipt):
        from docx import Document

        context = self._context(order, receipt)
        document = Document()
        document.add_heading("HomeChef", level=0)
        document.add_paragraph("Comprobante formal de pago")
        document.add_paragraph(f"Nro: {context['receipt_number']}")
        document.add_paragraph(f"Emitido: {context['issued_at']}")
        document.add_paragraph(f"Pedido: {context['order_id']}")

        table = document.add_table(rows=0, cols=2)
        table.style = "Table Grid"
        detail_rows = [
            ("Cliente", context["client_name"]),
            ("Cocinero", context["chef_name"]),
            ("Modalidad", context["fulfillment_type"]),
            ("Metodo de pago", context["payment_method"]),
            ("Estado del pago", context["payment_status"]),
            ("Referencia externa", context["external_reference"]),
            ("Direccion", context["address"]),
        ]
        for label, value in detail_rows:
            row = table.add_row().cells
            row[0].text = label
            row[1].text = value

        document.add_paragraph("")
        items_table = document.add_table(rows=1, cols=4)
        items_table.style = "Table Grid"
        header = items_table.rows[0].cells
        header[0].text = "Plato"
        header[1].text = "Cantidad"
        header[2].text = "Unitario"
        header[3].text = "Subtotal"
        for item in context["items"]:
            row = items_table.add_row().cells
            row[0].text = item["dish_name"]
            row[1].text = str(item["quantity"])
            row[2].text = item["unit_price"]
            row[3].text = item["subtotal"]

        document.add_paragraph("")
        totals_table = document.add_table(rows=0, cols=2)
        totals_table.style = "Table Grid"
        for label, value in [
            ("Subtotal", context["subtotal"]),
            ("Delivery", context["delivery_fee"]),
            ("Servicio", context["service_fee"]),
            ("Descuento", context["discount_total"]),
            ("Total", context["total"]),
        ]:
            row = totals_table.add_row().cells
            row[0].text = label
            row[1].text = value

        document.add_paragraph(f"Estado final del pedido al emitir: {context['order_status']}")
        buffer = BytesIO()
        document.save(buffer)
        return buffer.getvalue()

    def _render_pdf(self, order: Order, receipt: OrderReceipt):
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.units import mm
        from reportlab.pdfgen import canvas

        context = self._context(order, receipt)
        buffer = BytesIO()
        pdf = canvas.Canvas(buffer, pagesize=A4)
        width, height = A4
        y = height - 20 * mm

        pdf.setTitle(f"Comprobante {context['receipt_number']}")
        pdf.setFillColor(colors.HexColor("#0f172a"))
        pdf.setFont("Helvetica-Bold", 20)
        pdf.drawString(20 * mm, y, "HomeChef")
        y -= 8 * mm
        pdf.setFont("Helvetica", 11)
        pdf.drawString(20 * mm, y, "Comprobante formal de pago")
        y -= 10 * mm

        head_rows = [
            f"Nro: {context['receipt_number']}",
            f"Emitido: {context['issued_at']}",
            f"Pedido: {context['order_id']}",
            f"Cliente: {context['client_name']}",
            f"Cocinero: {context['chef_name']}",
            f"Modalidad: {context['fulfillment_type']}",
            f"Metodo de pago: {context['payment_method']}",
            f"Estado del pago: {context['payment_status']}",
            f"Referencia externa: {context['external_reference']}",
        ]
        for row in head_rows:
            pdf.drawString(20 * mm, y, row)
            y -= 6 * mm

        y -= 2 * mm
        pdf.setFont("Helvetica-Bold", 11)
        pdf.drawString(20 * mm, y, "Plato")
        pdf.drawString(95 * mm, y, "Cant.")
        pdf.drawString(120 * mm, y, "Unitario")
        pdf.drawString(155 * mm, y, "Subtotal")
        y -= 4 * mm
        pdf.line(20 * mm, y, 190 * mm, y)
        y -= 6 * mm
        pdf.setFont("Helvetica", 10)
        for item in context["items"]:
            pdf.drawString(20 * mm, y, item["dish_name"][:40])
            pdf.drawRightString(108 * mm, y, str(item["quantity"]))
            pdf.drawRightString(148 * mm, y, item["unit_price"])
            pdf.drawRightString(188 * mm, y, item["subtotal"])
            y -= 6 * mm
            if y < 40 * mm:
                pdf.showPage()
                y = height - 20 * mm
                pdf.setFont("Helvetica", 10)

        y -= 4 * mm
        pdf.line(120 * mm, y, 190 * mm, y)
        y -= 8 * mm
        total_rows = [
            ("Subtotal", context["subtotal"]),
            ("Delivery", context["delivery_fee"]),
            ("Servicio", context["service_fee"]),
            ("Descuento", context["discount_total"]),
            ("Total", context["total"]),
        ]
        for label, value in total_rows:
            pdf.drawString(120 * mm, y, label)
            pdf.drawRightString(188 * mm, y, value)
            y -= 6 * mm
        y -= 6 * mm
        pdf.drawString(20 * mm, y, f"Direccion: {context['address'][:95]}")
        y -= 6 * mm
        pdf.drawString(20 * mm, y, f"Estado final del pedido al emitir: {context['order_status']}")
        pdf.save()
        return buffer.getvalue()

    def _context(self, order: Order, receipt: OrderReceipt):
        client_name = self._profile_name(order.client)
        chef_name = self._chef_name(order.chef)
        address = getattr(order, "address", None)
        return {
            "receipt_number": receipt.receipt_number,
            "issued_at": timezone.localtime(receipt.issued_at).strftime("%Y-%m-%d %H:%M:%S"),
            "order_id": order.id,
            "client_name": client_name,
            "chef_name": chef_name,
            "fulfillment_type": "Delivery" if order.fulfillment_type == Order.FulfillmentType.DELIVERY else "Retiro",
            "payment_method": self._payment_method_label(receipt.payment_method),
            "payment_status": self._payment_status_label(receipt.payment_status),
            "order_status": self._order_status_label(receipt.order_status),
            "external_reference": receipt.external_reference or "-",
            "address": self._address_label(address),
            "subtotal": self._money(receipt.currency, receipt.subtotal),
            "delivery_fee": self._money(receipt.currency, receipt.delivery_fee),
            "service_fee": self._money(receipt.currency, receipt.service_fee),
            "discount_total": self._money(receipt.currency, receipt.discount_total),
            "total": self._money(receipt.currency, receipt.total),
            "items": [
                {
                    "dish_name": item.dish_name_snapshot,
                    "quantity": item.quantity,
                    "unit_price": self._money(receipt.currency, item.unit_price),
                    "subtotal": self._money(receipt.currency, item.subtotal),
                }
                for item in order.items.all()
            ],
        }

    def _money(self, currency: str, value):
        return f"{currency} {float(value):.2f}"

    def _address_label(self, address):
        if not address:
            return "Retiro en punto del cocinero"
        reference = f" | {address.reference}" if address.reference else ""
        return f"{address.line_1}{reference}"

    def _payment_method_label(self, value: str):
        return {
            Order.PaymentMethod.CASH: "Efectivo",
            Order.PaymentMethod.STRIPE_TEST: "Stripe test",
            Order.PaymentMethod.BITCOIN_COINGATE: "Bitcoin CoinGate",
            Order.PaymentMethod.QR_SIMULATED: "QR simulado",
        }.get(value, value)

    def _payment_status_label(self, value: str):
        return {
            OrderPayment.Status.CONFIRMED: "Confirmado",
            OrderPayment.Status.PENDING: "Pendiente",
            OrderPayment.Status.PROCESSING: "Procesando",
            OrderPayment.Status.FAILED: "Fallido",
            OrderPayment.Status.CANCELLED: "Cancelado",
            OrderPayment.Status.EXPIRED: "Expirado",
        }.get(value, value)

    def _order_status_label(self, value: str):
        return {
            Order.Status.AWAITING_CHEF_CONFIRMATION: "Esperando confirmacion",
            Order.Status.ACCEPTED: "Aceptado",
            Order.Status.PREPARING: "En preparacion",
            Order.Status.READY_FOR_PICKUP: "Listo para retiro",
            Order.Status.READY_FOR_DELIVERY: "Listo para delivery",
            Order.Status.OUT_FOR_DELIVERY: "En camino",
            Order.Status.DELIVERED: "Entregado",
            Order.Status.PICKED_UP: "Retirado",
            Order.Status.PAID: "Pagado",
            Order.Status.CANCELLED: "Cancelado",
            Order.Status.EXPIRED: "Expirado",
        }.get(value, value)

    def _get_order_for_client(self, client: UserProfile, order_id: str):
        order = (
            Order.objects.filter(id=str(order_id), client=client)
            .select_related("client", "chef", "address")
            .prefetch_related("items", "receipts")
            .first()
        )
        if not order:
            raise OrderReceiptServiceError("Pedido no encontrado para el cliente.", "order_not_found")
        return order

    def _get_order_for_chef(self, chef: UserProfile, order_id: str):
        order = (
            Order.objects.filter(id=str(order_id), chef=chef)
            .select_related("client", "chef", "address")
            .prefetch_related("items", "receipts")
            .first()
        )
        if not order:
            raise OrderReceiptServiceError("Pedido no encontrado para el cocinero.", "order_not_found")
        return order

    def _require_profile(self, user_id: str, role: str, code: str, message: str):
        parsed = self._parse_uuid(user_id)
        profile = UserProfile.objects.filter(supabase_user_id=parsed, role=role).first() if parsed else None
        if not profile:
            raise OrderReceiptServiceError(message, code)
        return profile

    def _parse_uuid(self, value):
        try:
            return UUID(str(value))
        except (TypeError, ValueError):
            return None

    def _chef_name(self, chef: UserProfile):
        profile = ChefProfile.objects.filter(user=chef).first() or ChefProfile.objects.filter(chef=chef).first()
        if profile and getattr(profile, "business_name", ""):
            return profile.business_name
        return self._profile_name(chef)

    def _profile_name(self, profile: UserProfile):
        name = f"{profile.first_name} {profile.last_name}".strip()
        return name or profile.full_name or profile.email
